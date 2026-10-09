# =============================================================================
# export_list.py
# Metadata & MediaClinic — Movie List Export (pure logic, no tkinter)
# Version: 0.18.3                                              ### NEW v0.18.3 ###
# Author:  Luiz Junqueira & Claude AI
#
# PURPOSE
# -------
# Turn scan-result row dicts into a file: CSV (comma or semicolon), tab-
# separated text, a plain text list, or an Excel workbook (openpyxl).
# Field selection and presets live here so the dialog in the main script is
# UI only and this module can be unit-tested headlessly.
#
# RULES
# -----
#   • NO tkinter import.  Pure functions only.
#   • CSV/TSV always go through the csv module, so commas, quotes and newlines
#     inside movie names are quoted correctly and never break the file.
#   • CSV is written with a UTF-8 BOM so Excel detects accented characters.
# =============================================================================

import csv
import os

# Status symbols (must match STATUS_* in the main script)
_STATUS_TEXT = {"⬤": "OK", "◐": "Warning", "○": "Missing", "✕": "Error"}


def _status_text(sym):
    return _STATUS_TEXT.get(sym, str(sym) if sym else "")


def _img_cell(row, key):
    if not row.get(f"{key}_exists"):
        return "Missing"
    st = _status_text(row.get(f"{key}_status"))
    q  = row.get(f"{key}_quality", "—")
    return f"{st} {q}".strip() if q and q != "—" else st


def _file_status(row, key):
    if not row.get(f"{key}_exists"):
        return "Missing"
    return "Error" if row.get(f"{key}_errors") else "OK"


def _errors_text(row, key):
    return "; ".join(f"L{e.get('line', 0)}: {e.get('message', '')}"
                     for e in row.get(f"{key}_errors", []) or [])


def _audio_text(row):
    tracks = sorted(row.get("audio_tracks", []) or [],
                    key=lambda t: t.get("language", ""))
    return ", ".join(f"{t.get('language', '')} ({t.get('channels')})"
                     if t.get("channels") else t.get("language", "")
                     for t in tracks)


def _votes_text(row):
    v = row.get("votes_int", 0) or 0
    return str(int(v)) if v else ""


# ── Field registry: (key, column header, getter(row) -> str) ─────────────────
EXPORT_FIELDS = [
    ("movie_name",    "Movie Name",     lambda r: r.get("movie_name") or r.get("subfolder", "")),
    ("year",          "Year",           lambda r: "" if r.get("movie_year", "-") in ("-", None) else str(r.get("movie_year"))),
    ("genres",        "Genres",         lambda r: "" if r.get("genre_display", "—") == "—" else r.get("genre_display", "")),
    ("rating",        "Rating",         lambda r: "" if r.get("rating_str", "-") == "-" else r.get("rating_str", "")),
    ("votes",         "Votes",          _votes_text),
    ("imdb_id",       "IMDB ID",        lambda r: r.get("source_imdb_val") or ""),
    ("tmdb_id",       "TMDB ID",        lambda r: r.get("source_tmdb_val") or ""),
    ("poster",        "Poster",         lambda r: _img_cell(r, "poster")),
    ("folder_img",    "Folder",         lambda r: _img_cell(r, "folder")),
    ("fanart",        "Fanart",         lambda r: _img_cell(r, "fanart")),
    ("backdrops",     "Backdrops",      lambda r: str(r.get("backdrop_count", 0))),
    ("nfo",           "NFO",            lambda r: _file_status(r, "nfo")),
    ("nfo_errors",    "NFO Errors",     lambda r: _errors_text(r, "nfo")),
    ("xml",           "XML",            lambda r: _file_status(r, "xml")),
    ("xml_errors",    "XML Errors",     lambda r: _errors_text(r, "xml")),
    ("language",      "Language",       lambda r: "" if r.get("language", "—") == "—" else r.get("language", "")),
    ("video_ext",     "Video",          lambda r: "" if r.get("video_count", 0) == 0 else r.get("video_ext", "")),
    ("video_size",    "Video Size",     lambda r: "" if r.get("video_count", 0) == 0 else r.get("video_size", "")),
    ("video_quality", "Video Quality",  lambda r: "" if r.get("video_quality", "—") in ("—", "-") else r.get("video_quality", "")),
    ("audio",         "Audio",          _audio_text),
    ("lang_ok",       "Lang OK?",       lambda r: "" if r.get("lang_ok", "—") == "—" else r.get("lang_ok", "")),
    ("subtitles",     "Subtitles",      lambda r: "" if r.get("subs_summary", "—") == "—" else r.get("subs_summary", "")),
    ("health",        "Health",         lambda r: (r.get("health_status") or "").capitalize()),
    ("health_reasons","Health Reasons", lambda r: "; ".join(r.get("health_reasons", []) or [])),
    ("folder_name",   "Folder Name",    lambda r: r.get("subfolder", "")),
    ("path",          "Path",           lambda r: r.get("subfolder_path", "")),
]

_FIELD_MAP = {k: (label, fn) for k, label, fn in EXPORT_FIELDS}
ALL_FIELD_KEYS = [k for k, _, _ in EXPORT_FIELDS]

PRESETS = {
    "Basic":    ["movie_name", "year"],
    "Standard": ["movie_name", "year", "genres", "rating", "votes",
                 "video_quality", "health"],
    "All":      list(ALL_FIELD_KEYS),
}

# (format id, label shown in dialog, default extension)
FORMATS = [
    ("csv_comma",     "CSV — comma separated (default)",                 ".csv"),
    ("csv_semicolon", "CSV — semicolon separated (Excel in CH/DE/PT/FR locales)", ".csv"),
    ("tsv",           "Tab-separated text (.txt)",                       ".txt"),
    ("txt",           "Plain text list — one movie per line (.txt)",      ".txt"),
    ("xlsx",          "Excel workbook (.xlsx)",                          ".xlsx"),
]
_FORMAT_EXT = {f: ext for f, _, ext in FORMATS}


def xlsx_available():
    """True when openpyxl can be imported (optional dependency)."""
    try:
        import openpyxl  # noqa: F401
        return True
    except Exception:
        return False


def default_extension(fmt):
    return _FORMAT_EXT.get(fmt, ".txt")


def build_table(rows, field_keys):
    """Return (headers, data) where data is a list of lists of strings."""
    keys = [k for k in field_keys if k in _FIELD_MAP] or ["movie_name"]
    headers = [_FIELD_MAP[k][0] for k in keys]
    data = []
    for r in rows:
        line = []
        for k in keys:
            try:
                v = _FIELD_MAP[k][1](r)
            except Exception:
                v = ""
            line.append("" if v is None else str(v))
        data.append(line)
    return headers, data


def _txt_line(keys, values):
    """Plain-text rendering: 'Name (Year)' when both present, else ' | ' join."""
    pairs = dict(zip(keys, values))
    name = pairs.get("movie_name", "")
    year = pairs.get("year", "")
    rest = [v for k, v in zip(keys, values)
            if k not in ("movie_name", "year") and v]
    head = f"{name} ({year})" if name and year else (name or year)
    return " | ".join([head] + rest) if head else " | ".join(rest)


def export_rows(rows, field_keys, fmt, path):
    """
    Write *rows* (list of scan-result dicts) to *path* in format *fmt*.
    Returns the number of movies written.  Raises on I/O error or when
    fmt == 'xlsx' and openpyxl is unavailable.
    """
    keys = [k for k in field_keys if k in _FIELD_MAP] or ["movie_name"]
    headers, data = build_table(rows, keys)

    if fmt in ("csv_comma", "csv_semicolon", "tsv"):
        delim = {"csv_comma": ",", "csv_semicolon": ";", "tsv": "\t"}[fmt]
        enc   = "utf-8-sig" if fmt != "tsv" else "utf-8"
        with open(path, "w", newline="", encoding=enc) as f:
            w = csv.writer(f, delimiter=delim, quotechar='"',
                           quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
            w.writerow(headers)
            w.writerows(data)
        return len(data)

    if fmt == "txt":
        with open(path, "w", encoding="utf-8") as f:
            for values in data:
                f.write(_txt_line(keys, values) + "\n")
        return len(data)

    if fmt == "xlsx":
        import openpyxl
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Movies"
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True)
        for values in data:
            out = []
            for k, v in zip(keys, values):
                if k in ("year", "votes", "backdrops") and v.isdigit():
                    out.append(int(v))
                elif k == "rating":
                    try:
                        out.append(float(v) if v else "")
                    except ValueError:
                        out.append(v)
                else:
                    out.append(v)
            ws.append(out)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, h in enumerate(headers, 1):
            width = max([len(h)] + [len(str(row[i - 1])) for row in data[:500]])
            ws.column_dimensions[get_column_letter(i)].width = min(max(8, width + 2), 60)
        wb.save(path)
        return len(data)

    raise ValueError(f"Unknown export format: {fmt}")


def suggested_filename(fmt, folder_name=""):
    base = "movie_list"
    if folder_name:
        safe = "".join(ch if ch.isalnum() or ch in "-_ " else "_" for ch in folder_name).strip()
        if safe:
            base = f"movie_list_{safe.replace(' ', '_')}"
    return base + default_extension(fmt)


if __name__ == "__main__":
    print("Fields:", ", ".join(ALL_FIELD_KEYS))
    print("xlsx available:", xlsx_available())
    print(os.path.basename(__file__), "is a library module — see selftest_v0_18_3.py")
